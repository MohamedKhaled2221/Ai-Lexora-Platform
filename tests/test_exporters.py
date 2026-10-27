from pathlib import Path

from docx import Document

from src.m1.docx_exporter import DocxExporter
from src.m1.pdf_exporter import PdfExporter
from src.m1.schemas import (
    AssembledDraft,
    AssembledSection,
    DraftStatus,
)


def make_draft():
    return AssembledDraft(
        draft_id="draft-test-001",
        request_id="request-test-001",
        template_id="template-test-001",
        template_version="1.0.0",
        document_type="nda",
        jurisdiction="EG",
        sections=[
            AssembledSection(
                section_id="sec-1",
                title="Test Section",
                text="Approved test content",
                order=1,
                source_type="template",
                source_id="template-test-001",
            ),
            AssembledSection(
                section_id="clause-1",
                title="Test Clause",
                text="Approved clause content",
                order=2,
                source_type="approved_clause",
                source_id="TEST-CLAUSE-001",
                clause_id="TEST-CLAUSE-001",
                clause_version="1.0.0",
            ),
        ],
        plain_text=(
            "Approved test content\n"
            "Approved clause content"
        ),
        status=DraftStatus.DRAFT,
    )


def test_docx_export(tmp_path):
    draft = make_draft()

    output_path = tmp_path / "output.docx"

    result = DocxExporter().export(
        draft,
        str(output_path),
    )

    assert result.success is True
    assert result.format == "docx"
    assert result.path == str(output_path)
    assert output_path.exists()

    # Re-open the generated DOCX and verify content.
    doc = Document(str(output_path))

    text = "\n".join(
        paragraph.text
        for paragraph in doc.paragraphs
    )

    assert "Approved test content" in text
    assert "Approved clause content" in text
    assert "[source:" not in text


def test_docx_export_creates_output_directory(tmp_path):
    draft = make_draft()

    output_path = (
        tmp_path
        / "nested"
        / "output"
        / "document.docx"
    )

    result = DocxExporter().export(
        draft,
        str(output_path),
    )

    assert result.success is True
    assert output_path.exists()


def test_pdf_export_missing_docx(tmp_path):
    missing_docx = tmp_path / "does_not_exist.docx"

    result = PdfExporter().export(
        str(missing_docx),
        str(tmp_path / "pdf"),
    )

    assert result.success is False
    assert result.format == "pdf"
    assert "does not exist" in result.error


def test_pdf_export_without_libreoffice(tmp_path, monkeypatch):
    source = tmp_path / "input.docx"

    doc = Document()
    doc.add_paragraph("Test document")
    doc.save(source)

    # Simulate LibreOffice not being installed.
    monkeypatch.setattr(
        "m1.pdf_exporter.shutil.which",
        lambda _: None,
    )

    result = PdfExporter().export(
        str(source),
        str(tmp_path / "pdf"),
    )

    assert result.success is False
    assert result.format == "pdf"
    assert "LibreOffice" in result.error