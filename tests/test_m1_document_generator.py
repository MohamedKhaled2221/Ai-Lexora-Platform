from __future__ import annotations

from datetime import (
    datetime,
    timedelta,
    timezone,
)
from pathlib import Path

import pytest
from docx import Document

from src.document_generator.errors import (
    ValidationError,
)
from src.document_generator.schemas import (
    FieldDefinition,
    GuidedDocumentRequest,
    TemplateDefinition,
)
from src.document_generator.service import (
    M1DocumentService,
)
from src.document_generator.template_registry import (
    TemplateRegistry,
)
from src.schemas.m5_contract import (
    ApprovalStatus,
    ApprovedClause,
    ClauseMetadata,
    ClauseType,
    DocumentType,
    Jurisdiction,
    TemplateClauseRequirement,
)


def _write_template(
    path: Path,
    *paragraphs: str,
) -> None:

    doc = Document()

    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)

    doc.save(path)


def _clause(
    clause_id: str,
    *,
    insertion_key: str,
    text: str,
) -> ApprovedClause:

    now = datetime.now(timezone.utc)

    return ApprovedClause(
        clause_id=clause_id,
        type=ClauseType.TERMINATION,
        jurisdiction=Jurisdiction.EG,
        document_types=[
            DocumentType.EMPLOYMENT_CONTRACT
        ],
        version="1.0.0",
        approval_status=(
            ApprovalStatus.APPROVED
        ),
        effective_date=(
            now - timedelta(days=1)
        ),
        owner="legal",
        source="approved-source",
        text_ar="نص قانوني معتمد",
        text_en=text,
        metadata=ClauseMetadata(
            insertion_key=insertion_key
        ),
    )


def _service(
    template_path: Path,
) -> M1DocumentService:

    template = TemplateDefinition(
        template_id="TMPL-EMP-EG-001",
        version="1.0.0",
        document_type=(
            DocumentType
            .EMPLOYMENT_CONTRACT
            .value
        ),
        jurisdiction=Jurisdiction.EG.value,
        file_path=str(template_path),
        fields=[
            FieldDefinition(
                name="employee_name",
                label="Employee",
            )
        ],
    )

    return M1DocumentService(
        TemplateRegistry([template])
    )


def _request() -> GuidedDocumentRequest:

    return GuidedDocumentRequest(
        request_id="REQ-001",
        document_type=(
            DocumentType
            .EMPLOYMENT_CONTRACT
            .value
        ),
        jurisdiction=Jurisdiction.EG.value,
        fields={
            "employee_name": "Alice"
        },
    )


def _requirement(
) -> TemplateClauseRequirement:

    return TemplateClauseRequirement(
        template_id="TMPL-EMP-EG-001",
        document_type=(
            DocumentType.EMPLOYMENT_CONTRACT
        ),
        jurisdiction=Jurisdiction.EG,
        required_clauses=[
            "CLZ-TERM-EG-001"
        ],
        optional_clauses=[
            "CLZ-CONF-EMP-EG-001"
        ],
        section_placeholders={
            "section_termination":
                "{{TERMINATION_CLAUSE}}",
            "section_confidentiality":
                "{{CONFIDENTIALITY_CLAUSE}}",
        },
        template_structure=(
            "{{TERMINATION_CLAUSE}}\n"
            "{{CONFIDENTIALITY_CLAUSE}}"
        ),
    )


def test_m1_uses_authoritative_m5_clause_contract(
    tmp_path: Path,
):

    template_path = tmp_path / "template.docx"

    _write_template(
        template_path,
        "Employee: {{employee_name}}",
        "Termination: "
        "{{TERMINATION_CLAUSE}}",
        "Optional: "
        "{{CONFIDENTIALITY_CLAUSE}}",
    )

    clause = _clause(
        "CLZ-TERM-EG-001",
        insertion_key="section_termination",
        text="Approved termination text.",
    )

    draft, manifest = (
        _service(template_path).build_draft(
            _request(),
            _requirement(),
            [clause],
        )
    )

    assert "Alice" in draft.plain_text

    assert (
        "Approved termination text."
        in draft.plain_text
    )

    assert (
        "{{TERMINATION_CLAUSE}}"
        not in draft.plain_text
    )

    assert (
        "{{CONFIDENTIALITY_CLAUSE}}"
        not in draft.plain_text
    )

    assert manifest.clause_refs == [
        {
            "clause_id":
                "CLZ-TERM-EG-001",
            "version": "1.0.0",
        }
    ]


def test_m1_accepts_m5_marker_as_insertion_key(
    tmp_path: Path,
):

    template_path = tmp_path / "template.docx"

    _write_template(
        template_path,
        "Employee: {{employee_name}}",
        "Termination: "
        "{{TERMINATION_CLAUSE}}",
        "Optional: "
        "{{CONFIDENTIALITY_CLAUSE}}",
    )

    clause = _clause(
        "CLZ-TERM-EG-001",
        insertion_key=(
            "{{TERMINATION_CLAUSE}}"
        ),
        text="Approved termination text.",
    )

    draft, _ = (
        _service(template_path).build_draft(
            _request(),
            _requirement(),
            [clause],
        )
    )

    assert (
        "Approved termination text."
        in draft.plain_text
    )


def test_m1_fails_when_selected_clause_marker_missing(
    tmp_path: Path,
):

    template_path = tmp_path / "template.docx"

    _write_template(
        template_path,
        "Employee: {{employee_name}}",
        "No legal marker is present here.",
    )

    clause = _clause(
        "CLZ-TERM-EG-001",
        insertion_key="section_termination",
        text="Approved termination text.",
    )

    with pytest.raises(
        ValidationError,
        match="was not found in M1 template",
    ):
        _service(template_path).build_draft(
            _request(),
            _requirement(),
            [clause],
        )


def test_m1_rejects_requirement_for_different_template(
    tmp_path: Path,
):

    template_path = tmp_path / "template.docx"

    _write_template(
        template_path,
        "Employee: {{employee_name}}",
        "Termination: "
        "{{TERMINATION_CLAUSE}}",
    )

    requirement = (
        _requirement().model_copy(
            update={
                "template_id":
                    "TMPL-OTHER"
            }
        )
    )

    with pytest.raises(
        ValueError,
        match="template_id mismatch",
    ):
        _service(template_path).build_draft(
            _request(),
            requirement,
            [
                _clause(
                    "CLZ-TERM-EG-001",
                    insertion_key=(
                        "section_termination"
                    ),
                    text=(
                        "Approved termination "
                        "text."
                    ),
                )
            ],
        )


def test_docx_export_contains_final_text_not_old_markers(
    tmp_path: Path,
):

    template_path = tmp_path / "template.docx"

    _write_template(
        template_path,
        "Employee: {{employee_name}}",
        "Termination: "
        "{{TERMINATION_CLAUSE}}",
        "Optional: "
        "{{CONFIDENTIALITY_CLAUSE}}",
    )

    service = _service(template_path)

    draft, _ = service.build_draft(
        _request(),
        _requirement(),
        [
            _clause(
                "CLZ-TERM-EG-001",
                insertion_key=(
                    "section_termination"
                ),
                text=(
                    "Approved termination text."
                ),
            )
        ],
    )

    output = tmp_path / "out.docx"

    result = service.docx_exporter.export(
        draft,
        str(output),
    )

    assert result.success is True

    exported = Document(str(output))

    text = "\n".join(
        paragraph.text
        for paragraph
        in exported.paragraphs
    )

    assert (
        "Approved termination text."
        in text
    )

    assert "{{TERMINATION_CLAUSE}}" not in text
    assert "{{CONFIDENTIALITY_CLAUSE}}" not in text
