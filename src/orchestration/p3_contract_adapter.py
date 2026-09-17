from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from src.document_generator.schemas import (
    AssembledDraft as M1AssembledDraft,
    GenerationManifest as M1GenerationManifest,
    GuidedDocumentRequest,
    TemplateDefinition,
)
from src.orchestration.grounding_validator import (
    P3GroundingContext,
    TrustedSourceText,
)
from src.schemas.m5_contract import (
    ApprovedClause,
    TemplateClauseRequirement,
)
from src.schemas.p3_generation_contract import (
    AssembledDraft,
    ClauseReference,
    DraftSection,
    DraftSourceType,
    GenerationManifest,
    SourceReference,
)


class P3ContractAdapterError(RuntimeError):
    """
    M1 -> shared Package 3 contract conversion failed.

    Fail closed rather than inventing provenance.
    """


@dataclass(frozen=True)
class P3ContractAdapterResult:
    draft: AssembledDraft
    manifest: GenerationManifest
    grounding_context: P3GroundingContext


class M1ToM4P3ContractAdapter:
    """
    Convert M1's deterministic output into the shared
    Package 3 M1 -> M4 contracts.

    M1's internal AssembledSection model is intentionally
    not copied directly because M1 stores paragraph-level
    assembly information differently from the shared M4
    provenance model.

    Instead, this adapter reconstructs provenance from the
    original M1 template, user fields, M5 placement rules,
    and exact approved M5 clause versions.
    """

    def adapt(
        self,
        *,
        request: GuidedDocumentRequest,
        template: TemplateDefinition,
        requirement: TemplateClauseRequirement,
        clauses: list[ApprovedClause],
        clause_language: str,
        m1_draft: M1AssembledDraft,
        m1_manifest: M1GenerationManifest,
    ) -> P3ContractAdapterResult:

        if clause_language not in {
            "ar",
            "en",
        }:
            raise P3ContractAdapterError(
                "clause_language must be 'ar' or 'en'"
            )

        self._validate_identity(
            request=request,
            template=template,
            requirement=requirement,
            m1_draft=m1_draft,
            m1_manifest=m1_manifest,
        )

        self._validate_manifest_clauses(
            clauses=clauses,
            m1_manifest=m1_manifest,
        )

        if m1_manifest.edits:
            raise P3ContractAdapterError(
                "M1 lawyer edits cannot be adapted until "
                "they use the shared structured lawyer-edit "
                "contract."
            )

        if m1_manifest.ai_transformation.used:
            raise P3ContractAdapterError(
                "M1 AI transformation metadata cannot be "
                "adapted safely because the current M1 "
                "metadata does not identify the exact "
                "target section and transformation ID."
            )

        template_path = Path(
            template.file_path or ""
        )

        if not template_path.is_file():
            raise P3ContractAdapterError(
                "M1 template file is unavailable: "
                f"{template.file_path!r}"
            )

        selected_clause_by_marker = (
            self._build_clause_marker_map(
                requirement=requirement,
                clauses=clauses,
            )
        )

        known_clause_markers = set(
            requirement
            .section_placeholders
            .values()
        )

        doc = Document(
            str(template_path)
        )

        shared_sections: list[
            DraftSection
        ] = []

        template_sources: list[
            TrustedSourceText
        ] = []

        section_number = 0

        for original_text in (
            self._iter_template_paragraphs(
                doc
            )
        ):
            section_number += 1

            section, sources = (
                self._materialize_section(
                    original_text=(
                        original_text
                    ),
                    section_number=(
                        section_number
                    ),
                    request=request,
                    template=template,
                    clauses_by_marker=(
                        selected_clause_by_marker
                    ),
                    known_clause_markers=(
                        known_clause_markers
                    ),
                    clause_language=(
                        clause_language
                    ),
                )
            )

            if section is None:
                continue

            shared_sections.append(
                section
            )

            template_sources.extend(
                sources
            )

        if not shared_sections:
            raise P3ContractAdapterError(
                "The M1 template produced no Package 3 "
                "sections."
            )

        reconstructed_plain_text = "\n".join(
            section.text
            for section in shared_sections
        )

        if (
            reconstructed_plain_text
            != m1_draft.plain_text
        ):
            raise P3ContractAdapterError(
                "M1 -> M4 adaptation did not reproduce "
                "the exact M1 assembled document text."
            )

        shared_draft = AssembledDraft(
            draft_id=m1_draft.draft_id,
            document_type=(
                m1_draft.document_type
            ),
            template_id=(
                m1_draft.template_id
            ),
            template_version=(
                m1_draft.template_version
            ),
            sections=shared_sections,
        )

        shared_manifest = GenerationManifest(
            manifest_id=(
                m1_manifest.generation_id
            ),
            draft_id=m1_draft.draft_id,
            template_id=(
                m1_draft.template_id
            ),
            template_version=(
                m1_draft.template_version
            ),
            clause_references=[
                ClauseReference(
                    clause_id=clause.clause_id,
                    clause_version=(
                        clause.version
                    ),
                )
                for clause in clauses
            ],
            user_fields={
                key: (
                    ""
                    if value is None
                    else str(value)
                )
                for key, value
                in request.fields.items()
            },
            lawyer_edit_references=[],
            ai_transformations=[],
            created_at=(
                self._manifest_created_at(
                    m1_manifest
                )
            ),
        )

        return P3ContractAdapterResult(
            draft=shared_draft,
            manifest=shared_manifest,
            grounding_context=(
                P3GroundingContext(
                    sources=template_sources
                )
            ),
        )

    # -----------------------------------------------------
    # Identity
    # -----------------------------------------------------

    @staticmethod
    def _validate_identity(
        *,
        request: GuidedDocumentRequest,
        template: TemplateDefinition,
        requirement: TemplateClauseRequirement,
        m1_draft: M1AssembledDraft,
        m1_manifest: M1GenerationManifest,
    ) -> None:

        if (
            template.template_id
            != requirement.template_id
        ):
            raise P3ContractAdapterError(
                "M1 template ID does not match "
                "the M5 template requirement."
            )

        if (
            request.document_type
            != requirement.document_type.value
        ):
            raise P3ContractAdapterError(
                "Request document type does not match "
                "the M5 template requirement."
            )

        if (
            request.jurisdiction
            != requirement.jurisdiction.value
        ):
            raise P3ContractAdapterError(
                "Request jurisdiction does not match "
                "the M5 template requirement."
            )

        if (
            m1_draft.draft_id
            != m1_manifest.draft_id
        ):
            raise P3ContractAdapterError(
                "M1 draft and manifest IDs differ."
            )

        if (
            m1_draft.template_id
            != template.template_id
        ):
            raise P3ContractAdapterError(
                "M1 draft template ID changed during "
                "assembly."
            )

        if (
            m1_manifest.template_id
            != template.template_id
        ):
            raise P3ContractAdapterError(
                "M1 manifest template ID differs from "
                "the selected template."
            )

    @staticmethod
    def _validate_manifest_clauses(
        *,
        clauses: list[ApprovedClause],
        m1_manifest: M1GenerationManifest,
    ) -> None:

        expected = {
            (
                clause.clause_id,
                clause.version,
            )
            for clause in clauses
        }

        actual = {
            (
                item.get("clause_id"),
                item.get("version"),
            )
            for item
            in m1_manifest.clause_refs
        }

        if actual != expected:
            raise P3ContractAdapterError(
                "M1 GenerationManifest clause references "
                "do not match the exact M5 clauses used "
                "for assembly."
            )

    # -----------------------------------------------------
    # Template traversal
    # -----------------------------------------------------

    @staticmethod
    def _iter_template_paragraphs(
        doc,
    ):
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                yield paragraph.text

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in (
                        cell.paragraphs
                    ):
                        if paragraph.text.strip():
                            yield paragraph.text

    # -----------------------------------------------------
    # Clause placement
    # -----------------------------------------------------

    def _build_clause_marker_map(
        self,
        *,
        requirement: TemplateClauseRequirement,
        clauses: list[ApprovedClause],
    ) -> dict[str, ApprovedClause]:

        marker_map: dict[
            str,
            ApprovedClause,
        ] = {}

        for clause in clauses:
            insertion_key = (
                clause.metadata.insertion_key
            )

            if not insertion_key:
                raise P3ContractAdapterError(
                    "Approved clause has no "
                    "insertion_key: "
                    + clause.clause_id
                )

            marker = (
                requirement
                .section_placeholders
                .get(insertion_key)
            )

            if marker is None:
                if (
                    insertion_key
                    in requirement
                    .section_placeholders
                    .values()
                ):
                    marker = insertion_key

                else:
                    raise P3ContractAdapterError(
                        "M5 clause placement is not "
                        "declared by the template "
                        "requirement: "
                        + clause.clause_id
                    )

            if marker in marker_map:
                raise P3ContractAdapterError(
                    "Multiple selected clauses target "
                    f"the same marker {marker!r}."
                )

            marker_map[marker] = clause

        return marker_map

    # -----------------------------------------------------
    # Section materialization
    # -----------------------------------------------------

    def _materialize_section(
        self,
        *,
        original_text: str,
        section_number: int,
        request: GuidedDocumentRequest,
        template: TemplateDefinition,
        clauses_by_marker: dict[
            str,
            ApprovedClause,
        ],
        known_clause_markers: set[str],
        clause_language: str,
    ) -> tuple[
        DraftSection | None,
        list[TrustedSourceText],
    ]:

        references: list[
            SourceReference
        ] = []

        sources: list[
            TrustedSourceText
        ] = []

        rendered_parts: list[str] = []

        pattern = re.compile(
            template.placeholder_pattern
        )

        cursor = 0
        template_part = 0

        for match in pattern.finditer(
            original_text
        ):
            static_text = (
                original_text[
                    cursor:match.start()
                ]
            )

            if static_text:
                rendered_parts.append(
                    static_text
                )

                if static_text.strip():
                    template_part += 1

                    source_id = (
                        f"{template.template_id}"
                        f"::section-{section_number}"
                        f"::part-{template_part}"
                    )

                    references.append(
                        SourceReference(
                            source_type=(
                                DraftSourceType
                                .TEMPLATE
                            ),
                            source_id=source_id,
                            source_version=(
                                template.version
                            ),
                        )
                    )

                    sources.append(
                        TrustedSourceText(
                            source_type=(
                                DraftSourceType
                                .TEMPLATE
                            ),
                            source_id=source_id,
                            source_version=(
                                template.version
                            ),
                            text=static_text,
                        )
                    )

            token = match.group(0)
            key = match.group(1)

            if (
                token
                in known_clause_markers
            ):
                clause = (
                    clauses_by_marker.get(
                        token
                    )
                )

                # Known but unselected optional
                # clause marker.
                if clause is None:
                    contribution = ""

                else:
                    contribution = (
                        self._clause_text(
                            clause,
                            clause_language,
                        )
                    )

                    references.append(
                        SourceReference(
                            source_type=(
                                DraftSourceType
                                .APPROVED_CLAUSE
                            ),
                            source_id=(
                                clause.clause_id
                            ),
                            source_version=(
                                clause.version
                            ),
                        )
                    )

                rendered_parts.append(
                    contribution
                )

            elif key in request.fields:
                value = request.fields[
                    key
                ]

                contribution = (
                    ""
                    if value is None
                    else str(value)
                )

                rendered_parts.append(
                    contribution
                )

                if contribution:
                    references.append(
                        SourceReference(
                            source_type=(
                                DraftSourceType
                                .USER_INPUT
                            ),
                            source_id=key,
                        )
                    )

            else:
                raise P3ContractAdapterError(
                    "Unresolved template placeholder "
                    f"{token!r} during M1 -> M4 "
                    "adaptation."
                )

            cursor = match.end()

        tail = original_text[cursor:]

        if tail:
            rendered_parts.append(
                tail
            )

            if tail.strip():
                template_part += 1

                source_id = (
                    f"{template.template_id}"
                    f"::section-{section_number}"
                    f"::part-{template_part}"
                )

                references.append(
                    SourceReference(
                        source_type=(
                            DraftSourceType.TEMPLATE
                        ),
                        source_id=source_id,
                        source_version=(
                            template.version
                        ),
                    )
                )

                sources.append(
                    TrustedSourceText(
                        source_type=(
                            DraftSourceType.TEMPLATE
                        ),
                        source_id=source_id,
                        source_version=(
                            template.version
                        ),
                        text=tail,
                    )
                )

        rendered_text = "".join(
            rendered_parts
        ).strip()

        if not rendered_text:
            return None, sources

        if not references:
            raise P3ContractAdapterError(
                "Materialized Package 3 section "
                "contains text but no trusted source."
            )

        return (
            DraftSection(
                section_id=(
                    f"section-{section_number}"
                ),
                title="",
                text=rendered_text,
                source_references=references,
            ),
            sources,
        )

    # -----------------------------------------------------
    # Exact clause text
    # -----------------------------------------------------

    @staticmethod
    def _clause_text(
        clause: ApprovedClause,
        language: str,
    ) -> str:

        text = (
            clause.text_ar
            if language == "ar"
            else clause.text_en
        )

        if not text.strip():
            raise P3ContractAdapterError(
                f"M5 clause {clause.clause_id!r} "
                f"has no {language!r} text."
            )

        return text.strip()

    # -----------------------------------------------------
    # Timestamps
    # -----------------------------------------------------

    @staticmethod
    def _manifest_created_at(
        manifest: M1GenerationManifest,
    ) -> datetime:

        raw = manifest.timestamps.get(
            "created_at"
        )

        if not raw:
            return datetime.now(
                timezone.utc
            )

        try:
            return datetime.fromisoformat(
                raw.replace(
                    "Z",
                    "+00:00",
                )
            )

        except ValueError as exc:
            raise P3ContractAdapterError(
                "M1 GenerationManifest contains "
                "an invalid created_at timestamp."
            ) from exc
