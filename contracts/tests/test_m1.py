from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document

from m1.document_assembler import DocumentAssembler
from m1.errors import ValidationError
from m1.schemas import (
    ApprovedClause,
    ClauseDecision,
    FieldDefinition,
    GuidedDocumentRequest,
    TemplateDefinition,
)


def make_template(tmp_path: Path) -> TemplateDefinition:
    path = tmp_path / "template.docx"

    doc = Document()
    doc.add_paragraph("Party A: {{party_a}}")
    doc.add_paragraph("Party B: {{party_b}}")
    doc.add_paragraph("Clause: {{CLAUSE:c1}}")
    doc.save(path)

    return TemplateDefinition(
        template_id="nda",
        version="1.0",
        document_type="NDA",
        jurisdiction="Egypt",
        file_path=str(path),
        fields=[
            FieldDefinition(
                name="party_a",
                label="Party A",
            ),
            FieldDefinition(
                name="party_b",
                label="Party B",
            ),
        ],
        required_clauses=["c1"],
        section_placeholders={
            "c1": "{{CLAUSE:c1}}",
        },
    )


def make_clause() -> ApprovedClause:
    return ApprovedClause(
        clause_id="c1",
        type="test",
        jurisdiction="Egypt",
        document_types=["NDA"],
        version="1.0",
        approval_status="approved",
        text_en="APPROVED TEST TEXT",
        insertion_key="c1",
        sequence=1,
        source="source-1",
    )


def make_request() -> GuidedDocumentRequest:
    return GuidedDocumentRequest(
        request_id=str(uuid4()),
        document_type="NDA",
        jurisdiction="Egypt",
        fields={
            "party_a": "A",
            "party_b": "B",
        },
    )


def test_successful_controlled_assembly(tmp_path):
    draft = DocumentAssembler().assemble(
        make_request(),
        make_template(tmp_path),
        [make_clause()],
        [
            ClauseDecision(
                clause_id="c1",
                decision="required",
            )
        ],
    )

    assert "A" in draft.plain_text
    assert "B" in draft.plain_text
    assert "APPROVED TEST TEXT" in draft.plain_text
    assert "{{" not in draft.plain_text


def test_missing_required_field(tmp_path):
    request = make_request().model_copy(
        update={
            "fields": {
                "party_a": "A",
            }
        }
    )

    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            request,
            make_template(tmp_path),
            [make_clause()],
        )


def test_incompatible_clause_blocked(tmp_path):
    bad = make_clause().model_copy(
        update={
            "jurisdiction": "UK",
        }
    )

    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            make_request(),
            make_template(tmp_path),
            [bad],
        )


def test_missing_required_clause(tmp_path):
    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            make_request(),
            make_template(tmp_path),
            [],
            [
                ClauseDecision(
                    clause_id="c1",
                    decision="required",
                )
            ],
        )


def test_special_condition_requires_lawyer_input(tmp_path):
    request = make_request().model_copy(
        update={
            "special_conditions": [
                "unsupported special condition"
            ]
        }
    )

    decision = ClauseDecision(
        clause_id="special",
        decision="requires_lawyer_input",
        reason="No approved variant exists",
    )

    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            request,
            make_template(tmp_path),
            [make_clause()],
            [decision],
        )


def test_invalid_insertion_key_blocked(tmp_path):
    bad = make_clause().model_copy(
        update={
            "insertion_key": "unknown_section",
        }
    )

    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            make_request(),
            make_template(tmp_path),
            [bad],
        )


def test_missing_insertion_key_blocked(tmp_path):
    bad = make_clause().model_copy(
        update={
            "insertion_key": None,
        }
    )

    with pytest.raises(ValidationError):
        DocumentAssembler().assemble(
            make_request(),
            make_template(tmp_path),
            [bad],
        )