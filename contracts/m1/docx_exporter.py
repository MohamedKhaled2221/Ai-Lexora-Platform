from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from .schemas import AssembledDraft, ExportResult


class DocxExporter:
    """
    Export an AssembledDraft to DOCX.

    The exporter does not generate or modify legal content.
    It only serializes the already assembled draft sections.
    """

    def export(
        self,
        draft: AssembledDraft,
        output_path: str,
    ) -> ExportResult:

        path = Path(output_path)
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        try:
            doc = Document()

            for section in sorted(
                draft.sections,
                key=lambda section: section.order,
            ):
                if section.title:
                    doc.add_heading(
                        section.title,
                        level=2,
                    )

                if section.text.strip():
                    doc.add_paragraph(section.text)

            doc.save(path)

            return ExportResult(
                success=True,
                format="docx",
                path=str(path),
                exported_at=datetime.now(
                    timezone.utc
                ).isoformat(),
            )

        except Exception as exc:

            return ExportResult(
                success=False,
                format="docx",
                error=str(exc),
                exported_at=datetime.now(
                    timezone.utc
                ).isoformat(),
            )