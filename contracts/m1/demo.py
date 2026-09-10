from __future__ import annotations

import tempfile
from pathlib import Path
from uuid import uuid4

from docx import Document

from .schemas import (
    ApprovedClause,
    ClauseDecision,
    FieldDefinition,
    GuidedDocumentRequest,
    TemplateDefinition,
)
from .service import M1DocumentService
from .template_registry import TemplateRegistry


def build_demo_template(path: Path) -> None:
    doc = Document()

    doc.add_heading("NON-LEGAL DEMO AGREEMENT", level=1)

    doc.add_paragraph("Party A: {{party_a_name}}")
    doc.add_paragraph("Party B: {{party_b_name}}")
    doc.add_paragraph("Effective Date: {{effective_date}}")

    doc.add_paragraph("Approved Section: {{CLAUSE:demo_clause_1}}")

    doc.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        # ---------------------------------------------------------
        # 1. Create a temporary DOCX template for the technical demo
        # ---------------------------------------------------------
        template_path = tmp / "demo.docx"
        build_demo_template(template_path)

        # ---------------------------------------------------------
        # 2. Define the template
        # ---------------------------------------------------------
        template = TemplateDefinition(
            template_id="demo",
            version="1.0",
            document_type="Demo Agreement",
            jurisdiction="Egypt",
            file_path=str(template_path),
            fields=[
                FieldDefinition(
                    name="party_a_name",
                    label="Party A",
                ),
                FieldDefinition(
                    name="party_b_name",
                    label="Party B",
                ),
                FieldDefinition(
                    name="effective_date",
                    label="Effective Date",
                ),
            ],
            section_placeholders={
                "demo_clause_1": "{{CLAUSE:demo_clause_1}}"
            },
        )

        # ---------------------------------------------------------
        # 3. Create a demo approved clause
        # ---------------------------------------------------------
        clause = ApprovedClause(
            clause_id="demo_clause_1",
            type="demo_section",
            jurisdiction="Egypt",
            document_types=["Demo Agreement"],
            version="1.0",
            approval_status="approved",
            source="DEMO-SOURCE",
            text_en=(
                "NON-LEGAL TEST TEXT — "
                "replace with an approved M5 clause."
            ),
            insertion_key="demo_clause_1",
            sequence=1,
        )

        # ---------------------------------------------------------
        # 4. Create the clause decision
        # ---------------------------------------------------------
        decision = ClauseDecision(
            clause_id="demo_clause_1",
            decision="required",
        )

        # ---------------------------------------------------------
        # 5. Create guided user input
        # ---------------------------------------------------------
        request = GuidedDocumentRequest(
            request_id=str(uuid4()),
            document_type="Demo Agreement",
            jurisdiction="Egypt",
            fields={
                "party_a_name": "Example A",
                "party_b_name": "Example B",
                "effective_date": "2026-09-07",
            },
        )

        # ---------------------------------------------------------
        # 6. Initialize M1 service
        # ---------------------------------------------------------
        service = M1DocumentService(
            TemplateRegistry([template])
        )

        # ---------------------------------------------------------
        # 7. Build the assembled draft + generation manifest
        # ---------------------------------------------------------
        draft, manifest = service.build_draft(
            request,
            [clause],
            [decision],
        )

        # ---------------------------------------------------------
        # 8. Export DOCX + PDF
        # ---------------------------------------------------------
        docx_result, pdf_result = service.export(
            draft,
            docx_path=str(tmp / "demo_output.docx"),
            pdf_dir=str(tmp / "exports"),
        )

        # ---------------------------------------------------------
        # 9. Print results
        # ---------------------------------------------------------
        print("Draft:", draft.draft_id)

        print("\nGeneration Manifest:")
        print(manifest.model_dump_json(indent=2))

        print("\nDOCX:")
        print(docx_result.model_dump())

        print("\nPDF:")
        print(
            pdf_result.model_dump()
            if pdf_result
            else None
        )


if __name__ == "__main__":
    main()